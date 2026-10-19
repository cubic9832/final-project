"""Extract official numeric household group statistics with explicit periods/units."""

import json
import re
from pathlib import Path

import openpyxl
from pypdf import PdfReader

ROOT = (Path(__file__).resolve().parents[3] / 'data/households')
# The first workbook sheet places the 2024/2025 survey columns side by side.
METRICS = {
    'household_share_pct': ('C', '%', '가구 비율'),
    'household_size_mean': ('E', '명', '평균 가구원 수'),
    'head_age_mean': ('G', '세', '평균 가구주 연령'),
    'assets_mean_won': ('I', '만원', '평균 총자산'),
    'financial_assets_mean_won': ('O', '만원', '평균 금융자산'),
    'savings_mean_won': ('U', '만원', '평균 저축액'),
    'deposit_savings_mean_won': ('AA', '만원', '평균 적립·예치식 저축'),
    'housing_deposit_mean_won': ('AG', '만원', '평균 전·월세보증금'),
    'real_assets_mean_won': ('AM', '만원', '평균 실물자산'),
    'debts_mean_won': ('BQ', '만원', '평균 총부채'),
    'financial_debts_mean_won': ('BW', '만원', '평균 금융부채'),
    'secured_loans_mean_won': ('CC', '만원', '평균 담보대출'),
    'unsecured_loans_mean_won': ('CI', '만원', '평균 신용대출'),
    'card_loans_mean_won': ('CO', '만원', '평균 신용카드 관련 대출'),
    'other_loans_mean_won': ('CU', '만원', '평균 기타 대출'),
    'rental_deposit_liabilities_mean_won': ('DA', '만원', '평균 임대보증금 부채'),
    'net_assets_mean_won': ('DG', '만원', '평균 순자산'),
    'annual_income_mean_won': ('DM', '만원', '평균 연간 가구소득'),
    'annual_nonconsumption_mean_won': ('DQ', '만원', '평균 연간 비소비지출'),
    'annual_disposable_income_mean_won': ('DU', '만원', '평균 연간 처분가능소득'),
    'annual_principal_interest_payment_mean_won': ('DY', '만원', '평균 연간 원리금상환액'),
}


def number(value):
    if isinstance(value, (int,float)):
        return float(value)
    if value in [None,'-','*']:
        return None
    return float(str(value).replace(',','').replace('p','').replace('*','').strip())


def extract_finance(source):
    book=openpyxl.load_workbook(ROOT/'sources'/source['filename'],data_only=True)
    sheet=book.worksheets[0]
    profiles=[];observations=[];dimension='전체'
    for r in range(10,48):
        label=sheet.cell(r,1).value
        if label:
            dimension=re.sub(r'\s+','',label)
        subgroup=sheet.cell(r,2).value
        group=re.sub(r'\s+','',subgroup) if subgroup else '전체'
        nested=group.startswith('·')
        for year,offset in [(2024,0),(2025,1)]:
            profile={'survey_year':year,'dimension':dimension,'group':group,
                'overlapping_subgroup':nested,'asset_debt_reference_date':f'{year}-03-31',
                'income_reference_year':year-1,'unit_of_observation':'가구 특성별 집계',
                'source_url':source['sourceUrl'],'source_sheet':sheet.title,'source_row':r,
                'checked_on':'2026-10-14'}
            for key,(column,unit,label) in METRICS.items():
                col=openpyxl.utils.column_index_from_string(column)+offset
                cell=sheet.cell(r,col)
                value=number(cell.value)
                profile[key]=None if value is None else (value*10000 if unit=='만원' else value)
                observations.append({'survey_year':year,'dimension':dimension,'group':group,
                    'overlapping_subgroup':nested,'metric':label,'metric_key':key,'population':'전체가구',
                    'reference_period':str(year-1) if key.startswith('annual_') else f'{year}-03-31',
                    'source_value':value,'source_unit':unit,
                    'normalized_value':profile[key],'normalized_unit':'원' if unit=='만원' else unit,
                    'source_url':source['sourceUrl'],'source_sheet':sheet.title,'source_cell':cell.coordinate,
                    'source_number_format':cell.number_format,'checked_on':'2026-10-14'})
            for key in ['income','nonconsumption','disposable_income','principal_interest_payment']:
                annual=profile[f'annual_{key}_mean_won']
                profile[f'monthly_{key}_mean_derived_won']=annual/12 if annual is not None else None
            profiles.append(profile)
    return profiles,observations


def parse_quintile_line(line):
    # The PDF extractor sometimes inserts spaces between digits of the first value.
    line=re.sub(r'(?<!\d)[1-4]\)', '',line)
    tokens=list(re.finditer(r'-?\d[\d,]*(?:\.\d+)?p?|(?<!\S)-(?!\S)',line))
    if len(tokens)<10:
        return None
    tail=tokens[-9:]
    prefix=line[:tail[0].start()].strip()
    first=re.search(r'-?\d',prefix)
    if not first:
        return None
    label=re.sub(r'\s+','',prefix[:first.start()]).replace('*','')
    first_value=re.sub(r'\s+','',prefix[first.start():])
    values=[number(first_value)]+[number(token.group()) for token in tail]
    return label,values[::2],values[1::2],('*' in prefix)


def extract_expenditure(source):
    reader=PdfReader(ROOT/'sources'/source['filename'])
    observations=[]
    for page_index,table_id,end in [(14,'<표 5>','주:'),(15,'<표 6>','<그림')]:
        text=reader.pages[page_index].extract_text()
        table=text.split(table_id,1)[1].split(end,1)[0]
        for line in table.splitlines():
            parsed=parse_quintile_line(line)
            if not parsed:
                continue
            label,values,changes,warning=parsed
            if '가구원수' in label: unit='명'
            elif '가구주연령' in label: unit='세'
            elif '%' in label: unit='%'
            else: unit='천원'
            for q,(value,change) in enumerate(zip(values,changes),1):
                observations.append({'period':'2025Q4','population':'전국 1인 이상 일반가구(농림어가 포함)',
                    'income_quintile':q,'metric':label,'source_value':value,'source_unit':unit,
                    'normalized_value':value*1000 if unit=='천원' and value is not None else value,
                    'normalized_unit':'원/월' if unit=='천원' else unit,
                    'year_on_year_change':change,'change_unit':'%p' if unit=='%' else '%',
                    'high_rse_warning':warning and q==1,'source_url':source['sourceUrl'],'source_page':page_index+1,
                    'source_table':table_id,'checked_on':'2026-10-14'})
    profiles=[]
    for q in range(1,6):
        row={'period':'2025Q4','income_quintile':q,'population':'전국 1인 이상 일반가구(농림어가 포함)'}
        for obs in observations:
            if obs['income_quintile']!=q: continue
            if obs['metric'] in row and row[obs['metric']]!=obs['normalized_value']:
                raise ValueError('Inconsistent duplicated metric '+obs['metric'])
            row[obs['metric']]=obs['normalized_value']
        row.update({'source_url':source['sourceUrl'],'source_pages':'15;16','checked_on':'2026-10-14'})
        profiles.append(row)
    return profiles,observations


def main():
    manifest=json.loads((ROOT/'source-manifest.json').read_text(encoding='utf-8'))
    financial_source=next(s for s in manifest if s['filename'].endswith('.xlsx'))
    expense_source=next(s for s in manifest if s['filename'].endswith('.pdf'))
    finance,finance_long=extract_finance(financial_source)
    expenditure,expenditure_long=extract_expenditure(expense_source)
    result={'financial_profiles':finance,'financial_metrics':finance_long,
            'expenditure_profiles':expenditure,'expenditure_metrics':expenditure_long}
    (ROOT/'official-household-statistics.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({key:len(value) for key,value in result.items()}))


if __name__=='__main__':
    main()
